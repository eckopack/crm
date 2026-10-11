"""
Management-Command: exportiert den aktuellen CRM-Stand (Firmen/Kontakte/Deals)
in dieselbe 3-Blatt-Struktur wie die ursprüngliche Lead-Recherche-Liste, zum
Abgleich nach einem Import und als Vorlage für die Weiterpflege.

Nutzung:
    python manage.py export_leadliste /pfad/zur/export.xlsx
"""
import openpyxl
from django.core.management.base import BaseCommand

from vertrieb.models import Deal, Firma, Kontakt


class Command(BaseCommand):
    help = "Exportiert Firmen/Kontakte/Deals im Format der Lead-Recherche-Liste (xlsx)."

    def add_arguments(self, parser):
        parser.add_argument("pfad", type=str, help="Zielpfad für die .xlsx-Datei")

    def handle(self, *args, **options):
        pfad = options["pfad"]
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        self._export_firmen(wb)
        self._export_kontakte(wb)
        self._export_leads(wb)

        wb.save(pfad)
        self.stdout.write(self.style.SUCCESS(
            f"Export geschrieben nach {pfad} "
            f"({Firma.objects.count()} Firmen, {Kontakt.objects.count()} Kontakte, {Deal.objects.count()} Deals)."
        ))

    def _header(self, ws, titel, spalten):
        ws.append([titel])
        ws.append(spalten)
        for cell in ws[2]:
            cell.font = openpyxl.styles.Font(bold=True)

    def _export_firmen(self, wb):
        ws = wb.create_sheet("Firmen")
        self._header(ws, "CRM-Export · Firmen", [
            "Firmen-ID (extern)", "CRM-ID", "Name", "Geschäftsbereich", "Status", "Segment/Branche",
            "Land", "Straße", "PLZ", "Ort", "Telefon", "E-Mail", "Website", "Zuständiger", "Notizen",
        ])
        for f in Firma.objects.all().order_by("name"):
            ws.append([
                f.externe_id, f.pk, f.name, f.get_geschaeftsbereich_display(), f.get_status_display(),
                f.branche_kunde, f.land, f.strasse, f.plz, f.ort, f.telefon, f.email, f.website,
                str(f.zustaendiger) if f.zustaendiger else "", f.notizen,
            ])

    def _export_kontakte(self, wb):
        ws = wb.create_sheet("Kontakte")
        self._header(ws, "CRM-Export · Kontakte", [
            "Kontakt-ID (extern)", "CRM-ID", "Firma", "Vorname", "Nachname", "Position", "E-Mail",
            "Telefon", "Mobil", "Hauptansprechpartner", "Notizen",
        ])
        for k in Kontakt.objects.select_related("firma").order_by("firma__name", "nachname"):
            ws.append([
                k.externe_id, k.pk, k.firma.name, k.vorname, k.nachname, k.position, k.email,
                k.telefon, k.mobil, "Ja" if k.ist_hauptansprechpartner else "Nein", k.notizen,
            ])

    def _export_leads(self, wb):
        ws = wb.create_sheet("Leads")
        self._header(ws, "CRM-Export · Deals/Leads", [
            "Lead-ID (extern)", "CRM-ID", "Titel", "Firma", "Kontakt", "Geschäftsbereich", "Phase",
            "Wert (EUR)", "Potenzial (t)", "Wahrscheinlichkeit (%)", "Gewichteter Wert (EUR)",
            "Erw. Abschluss", "Zuständiger", "Notizen",
        ])
        for d in Deal.objects.select_related("firma", "kontakt").order_by("firma__name", "titel"):
            ws.append([
                d.externe_id, d.pk, d.titel, d.firma.name, str(d.kontakt) if d.kontakt else "",
                d.get_geschaeftsbereich_display(), d.get_phase_display(), float(d.wert),
                float(d.potenzial_tonnen) if d.potenzial_tonnen is not None else None,
                d.wahrscheinlichkeit, float(d.gewichteter_wert),
                d.erwartetes_abschlussdatum.isoformat() if d.erwartetes_abschlussdatum else "",
                str(d.zustaendiger) if d.zustaendiger else "", d.notizen,
            ])
