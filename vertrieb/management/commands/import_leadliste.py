"""
Management-Command zum Import der ECKO-PACK Lead-Recherche-Excel-Liste
(Blätter: Firmen, Kontakte, Leads) in das CRM.

Nutzung:
    python manage.py import_leadliste /pfad/zur/Leadliste.xlsx
    python manage.py import_leadliste /pfad/zur/Leadliste.xlsx --dry-run

Der Import ist idempotent: Firmen/Kontakte/Deals werden über das Feld
"externe_id" (Firmen-ID / Kontakt-ID / Lead-ID aus der Excel-Liste) erkannt
und bei einem erneuten Import aktualisiert statt dupliziert. So kann diese
Datei als lebende "Masterdatei" für neue Leads weiterverwendet werden -
einfach neue Zeilen ergänzen und das Command erneut laufen lassen.

Annahmen (siehe Chat mit Björn vom 11.10.2026):
- "Potenzial (t)" ist eine Tonnen-Mengenangabe, komplett getrennt vom
  EUR-Wertfeld (Deal.wert). Deal.wert bleibt 0, bis echte Preise vorliegen.
- Alle importierten Firmen/Deals bekommen Geschäftsbereich
  "Einweg-/Serviceprodukte" (einweg_serviceprodukte).
- Firma.name = volle Firmierung; Marke/Kurzname wird zusätzlich in die
  Notizen geschrieben.
- Firma.status = "lead" (alles Neukunden-Akquise, noch kein Bestandskunde).
- "Nächster Schritt" wird als offene Aktivität (Aufgabe) angelegt, Datum =
  Importzeitpunkt, da "Wiedervorlage" in der Quelldatei durchgehend leer ist.
- Keine Zuständigkeit wird vergeben (Spalte ist in der Quelle für alle 36
  Leads leer) - bleibt im CRM offen zur manuellen Zuordnung.
"""
from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from vertrieb.models import Aktivitaet, Deal, Firma, Kontakt

GESCHAEFTSBEREICH_IMPORT = "einweg_serviceprodukte"


def _s(value):
    """Wert zu getrimmtem String, None/leer -> ''."""
    if value is None:
        return ""
    return str(value).strip()


def _dec(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


class Command(BaseCommand):
    help = "Importiert Firmen/Kontakte/Leads aus der ECKO-PACK Lead-Recherche-Excel-Liste."

    def add_arguments(self, parser):
        parser.add_argument("pfad", type=str, help="Pfad zur .xlsx-Datei")
        parser.add_argument(
            "--dry-run", action="store_true",
            help="Nur anzeigen, was passieren würde - ohne etwas zu speichern.",
        )

    def handle(self, *args, **options):
        pfad = options["pfad"]
        dry_run = options["dry_run"]

        try:
            wb = openpyxl.load_workbook(pfad, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Datei nicht gefunden: {pfad}")

        for sheet in ("Firmen", "Kontakte", "Leads"):
            if sheet not in wb.sheetnames:
                raise CommandError(f"Blatt '{sheet}' fehlt in der Datei - Import abgebrochen.")

        stats = {
            "firmen_neu": 0, "firmen_aktualisiert": 0,
            "kontakte_neu": 0, "kontakte_aktualisiert": 0,
            "deals_neu": 0, "deals_aktualisiert": 0,
            "aktivitaeten_neu": 0,
            "warnungen": [],
        }

        with transaction.atomic():
            firma_map = self._import_firmen(wb["Firmen"], stats)
            kontakt_map, lead_to_kontakte = self._import_kontakte(wb["Kontakte"], firma_map, stats)
            self._import_leads(wb["Leads"], firma_map, lead_to_kontakte, stats)

            if dry_run:
                transaction.set_rollback(True)
                self.stdout.write(self.style.WARNING("--dry-run: Es wurde NICHTS gespeichert."))

        self.stdout.write(self.style.SUCCESS(
            f"Firmen: {stats['firmen_neu']} neu, {stats['firmen_aktualisiert']} aktualisiert\n"
            f"Kontakte: {stats['kontakte_neu']} neu, {stats['kontakte_aktualisiert']} aktualisiert\n"
            f"Deals: {stats['deals_neu']} neu, {stats['deals_aktualisiert']} aktualisiert\n"
            f"Aktivitäten (Nächster Schritt): {stats['aktivitaeten_neu']} neu"
        ))
        if stats["warnungen"]:
            self.stdout.write(self.style.WARNING(f"\n{len(stats['warnungen'])} Warnung(en):"))
            for w in stats["warnungen"]:
                self.stdout.write(self.style.WARNING(f"  - {w}"))

    # ------------------------------------------------------------------
    def _import_firmen(self, ws, stats):
        rows = list(ws.iter_rows(min_row=10, values_only=True))
        firma_map = {}
        for r in rows:
            if not r[0]:
                continue
            firmen_id, marke, firmierung, segment, land, strasse, plz, ort, telefon, email, website, \
                datenstatus, einordnung, quelle, weitere_quellen, recherche_am = r[:16]

            name = _s(firmierung) or _s(marke)
            if not name:
                stats["warnungen"].append(f"Firma {firmen_id}: kein Name - übersprungen.")
                continue

            notiz_teile = []
            if marke and _s(marke) != name:
                notiz_teile.append(f"Marke/Kurzname: {_s(marke)}")
            if datenstatus:
                notiz_teile.append(f"Datenstatus: {_s(datenstatus)}")
            if einordnung:
                notiz_teile.append(f"Prüfhinweis: {_s(einordnung)}")
            if quelle:
                notiz_teile.append(f"Quelle: {_s(quelle)}")
            if weitere_quellen:
                notiz_teile.append(f"Weitere Quellen: {_s(weitere_quellen)}")
            if recherche_am:
                datum = recherche_am.strftime("%d.%m.%Y") if hasattr(recherche_am, "strftime") else _s(recherche_am)
                notiz_teile.append(f"Recherche am: {datum}")
            notizen = "\n".join(notiz_teile)

            defaults = {
                "name": name,
                "geschaeftsbereich": GESCHAEFTSBEREICH_IMPORT,
                "branche_kunde": _s(segment)[:150],
                "website": _s(website),
                "telefon": _s(telefon)[:50],
                "email": _s(email),
                "strasse": _s(strasse)[:200],
                "plz": _s(plz)[:10],
                "ort": _s(ort)[:100],
                "land": _s(land) or "Deutschland",
                "notizen": notizen,
            }

            obj, created = Firma.objects.get_or_create(
                externe_id=_s(firmen_id), defaults={**defaults, "status": "lead"}
            )
            if created:
                stats["firmen_neu"] += 1
            else:
                for k, v in defaults.items():
                    setattr(obj, k, v)
                obj.save()
                stats["firmen_aktualisiert"] += 1

            if not land:
                stats["warnungen"].append(f"Firma {firmen_id} ({name}): kein Land angegeben.")
            firma_map[_s(firmen_id)] = obj
        return firma_map

    # ------------------------------------------------------------------
    def _import_kontakte(self, ws, firma_map, stats):
        rows = list(ws.iter_rows(min_row=10, values_only=True))
        kontakt_map = {}
        lead_to_kontakte = {}
        for r in rows:
            if not r[0]:
                continue
            kontakt_id, firmen_id, firmenbezug, name, funktion, email, durchwahl, mobil, \
                lead_bezug, pruefstatus, notiz, herkunft = r[:12]

            firma = firma_map.get(_s(firmen_id))
            if firma is None:
                stats["warnungen"].append(f"Kontakt {kontakt_id} ({name}): Firma {firmen_id} nicht gefunden.")
                continue

            teile = _s(name).split()
            if len(teile) >= 2:
                vorname, nachname = teile[0], " ".join(teile[1:])
            elif teile:
                vorname, nachname = "", teile[0]
                stats["warnungen"].append(f"Kontakt {kontakt_id}: Name '{name}' konnte nicht in Vor-/Nachname getrennt werden.")
            else:
                vorname, nachname = "", "(ohne Namen)"

            notiz_teile = []
            if pruefstatus:
                notiz_teile.append(f"Prüfstatus: {_s(pruefstatus)}")
            if notiz:
                notiz_teile.append(_s(notiz))
            if herkunft:
                notiz_teile.append(f"Herkunft: {_s(herkunft)}")
            notizen = "\n".join(notiz_teile)

            defaults = {
                "firma": firma,
                "vorname": vorname,
                "nachname": nachname,
                "position": _s(funktion)[:150],
                "email": _s(email),
                "telefon": _s(durchwahl)[:50],
                "mobil": _s(mobil)[:50],
                "notizen": notizen,
            }

            obj, created = Kontakt.objects.get_or_create(externe_id=_s(kontakt_id), defaults=defaults)
            if created:
                stats["kontakte_neu"] += 1
            else:
                for k, v in defaults.items():
                    setattr(obj, k, v)
                obj.save()
                stats["kontakte_aktualisiert"] += 1

            kontakt_map[_s(kontakt_id)] = obj
            for lead_id in [x.strip() for x in _s(lead_bezug).split("/") if x.strip()]:
                lead_to_kontakte.setdefault(lead_id, []).append(obj)

        # Hauptansprechpartner: wenn eine Firma genau einen importierten Kontakt hat,
        # automatisch als Hauptansprechpartner markieren.
        firma_kontakt_count = {}
        for k in kontakt_map.values():
            firma_kontakt_count[k.firma_id] = firma_kontakt_count.get(k.firma_id, 0) + 1
        for k in kontakt_map.values():
            if firma_kontakt_count[k.firma_id] == 1 and not k.ist_hauptansprechpartner:
                k.ist_hauptansprechpartner = True
                k.save(update_fields=["ist_hauptansprechpartner"])

        return kontakt_map, lead_to_kontakte

    # ------------------------------------------------------------------
    def _import_leads(self, ws, firma_map, lead_to_kontakte, stats):
        rows = list(ws.iter_rows(min_row=10, values_only=True))
        for r in rows:
            if not r[0]:
                continue
            lead_id, titel, firmen_id, produktbedarf, potenzial_t, status, prioritaet, zustaendig, \
                naechster_schritt, wiedervorlage, kontaktmonat, termin, herkunft, gaiasip, \
                pruefhinweis, gespraechsnotiz, originalzeile = r[:17]

            firma = firma_map.get(_s(firmen_id))
            if firma is None:
                stats["warnungen"].append(f"Lead {lead_id} ({titel}): Firma {firmen_id} nicht gefunden - übersprungen.")
                continue

            kontakte = lead_to_kontakte.get(_s(lead_id), [])
            kontakt = kontakte[0] if kontakte else None

            notiz_teile = []
            if produktbedarf:
                notiz_teile.append(f"Produktbedarf: {_s(produktbedarf)}")
            if gaiasip:
                notiz_teile.append(f"Gaiasip-Interesse: {_s(gaiasip)}")
            if herkunft:
                notiz_teile.append(f"Herkunft/Kürzel: {_s(herkunft)}")
            if kontaktmonat:
                monat = kontaktmonat.strftime("%m/%Y") if hasattr(kontaktmonat, "strftime") else _s(kontaktmonat)
                notiz_teile.append(f"Kontaktmonat (Vorlage): {monat}")
            if termin:
                notiz_teile.append(f"Termin (Vorlage): {_s(termin)}")
            if pruefhinweis:
                notiz_teile.append(f"Prüfhinweis: {_s(pruefhinweis)}")
            if gespraechsnotiz:
                notiz_teile.append(f"Gesprächsnotiz: {_s(gespraechsnotiz)}")
            notizen = "\n".join(notiz_teile)

            defaults = {
                "titel": _s(titel)[:200] or f"Lead {lead_id}",
                "firma": firma,
                "kontakt": kontakt,
                "geschaeftsbereich": GESCHAEFTSBEREICH_IMPORT,
                # 24.10.2026, auf Wunsch von Björn: Mengenpotenzial ist jetzt nach Trinkhalm/
                # Besteck aufgeteilt (siehe Deal-Model). Die Lead-Liste kennt diese Trennung
                # nicht - der importierte Wert landet deshalb vollständig bei "Besteck" (Björns
                # Entscheidung: "erstmal alle mengenpotenziale bei besteck lassen, ich ändere
                # danach manuell"), Trinkhalm bleibt leer und wird bei Bedarf manuell gepflegt.
                "potenzial_tonnen_besteck": _dec(potenzial_t),
                "notizen": notizen,
            }

            obj, created = Deal.objects.get_or_create(
                externe_id=_s(lead_id),
                defaults={**defaults, "phase": "lead", "wahrscheinlichkeit": 10},
            )
            if created:
                stats["deals_neu"] += 1
            else:
                for k, v in defaults.items():
                    setattr(obj, k, v)
                obj.save()
                stats["deals_aktualisiert"] += 1

            if naechster_schritt and created:
                Aktivitaet.objects.create(
                    typ="aufgabe",
                    betreff=_s(naechster_schritt)[:200],
                    beschreibung="Automatisch aus Lead-Recherche-Liste übernommen (Spalte 'Nächster Schritt').",
                    datum=timezone.now(),
                    firma=firma,
                    kontakt=kontakt,
                    deal=obj,
                )
                stats["aktivitaeten_neu"] += 1
