"""
Management-Command zum Import des Artikelstamms (Blatt "Artikel") aus der
ECKO-PACK-Artikel-Schnittstelle (Excel) in das CRM.

Nutzung:
    python manage.py import_artikelstamm /pfad/zur/Datei.xlsx
    python manage.py import_artikelstamm /pfad/zur/Datei.xlsx --dry-run

Der Import ist idempotent: Artikel werden über das Feld "Artikelnummer" erkannt und bei
einem erneuten Import aktualisiert statt dupliziert. Preisstaffeln eines Artikels werden
bei jedem Import komplett neu aus der Datei geschrieben (alte Staffeln dieses Artikels
werden vorher gelöscht) - so bleibt die Datei eine lebende Masterdatei: Zeile ändern,
neu importieren.

Stand 15.10.2026 (Björns Entscheidungen, siehe Doku-Dokument "LUAMO Artikelstamm -
Abgleich Excel ↔ CRM" und Chat-Verlauf):

- Format seit 15.10.: pro VK-Stufe gibt es zwei Spalten nebeneinander, "Mindestmenge i
  (Stk)" und "VK i (€/Stk)" - die Mindestmenge steht jetzt direkt und real in der Datei
  (warengruppenabhängig von Björn gepflegt), NICHT mehr als hart im Code hinterlegte
  Tabelle. Eine Preisstaffel wird nur angelegt, wenn BEIDE Zellen eines Paares gefüllt
  sind; leere Paare (z. B. Großbesteck/Kleinbesteck ohne 9. Stufe, oder Artikel ohne
  jede Mengenstaffel wie Muster/Verpackung-Marketing) werden einfach übersprungen.

- Spalte "Einspielfreigabe (Ja/Nein)" ist ab jetzt ein ECHTES Gate, kein reiner
  Excel-Hinweis mehr: nur Zeilen mit "Ja" (Groß-/Kleinschreibung egal) werden
  importiert. Alles andere (leer, "Nein", Tippfehler) wird übersprungen und separat
  gezählt. Vorher stand das nur als Konvention in der Anleitung, ohne Durchsetzung im
  Code - genau das sollte diese Spalte eigentlich schon immer tun.

- Bezeichnung/Beschreibung können zweizeilig sein (1. Zeile Deutsch, 2. Zeile Englisch,
  per Zeilenumbruch in derselben Zelle). Seit 16.10.2026 wird der Zeilenumbruch beim
  Import durch " // " ersetzt (siehe _mehrsprachig) statt unverändert übernommen zu
  werden - ein reiner \n verschwindet beim Anzeigen in einzeiligen <input>-Feldern und
  in der Artikel-Dropdown-Liste spurlos, sodass DE und EN ohne Trennung aneinanderkleben.

- Die Spalte "Geschäftsbereich *" in der Datei enthält KEINE der fünf ECKO-PACK-
  Unternehmensbereiche, sondern eine feinere Produktkategorie (z. B. "Einweg-/
  Serviceprodukte", "Großbesteck", "Kleinbesteck", "Muster", "Verpackung/Marketing").
  Dieser Text wird deshalb in das Feld Artikel.warengruppe geschrieben;
  Artikel.geschaeftsbereich (der Unternehmensbereich) wird für ALLE importierten Zeilen
  fest auf "einweg_serviceprodukte" gesetzt, da alle Kategorien dort einsortiert sind.

- Artikel ohne jeden Preis (weder Flatrate-Verkaufspreis noch eine Preisstaffel) werden
  TROTZDEM angelegt (Stammdaten/Lagerbestand sind auch ohne Preis nützlich), aber mit
  einer Warnung gemeldet.

- Status (Spalte "Status"): wird informativ in Artikel.status übernommen (keine Sperre
  für Angebote/Packscheine), unbekannte Texte fallen auf "in_vorbereitung" zurück.
"""
from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from vertrieb.models import Artikel, EINHEIT_CHOICES, Preisstaffel

GESCHAEFTSBEREICH_IMPORT = "einweg_serviceprodukte"

STATUS_MAPPING = {
    "in vorbereitung": "in_vorbereitung",
    "entwurf": "in_vorbereitung",
    "aktiv": "aktiv",
    "auslaufend": "auslaufend",
    "gesperrt": "gesperrt",
}

_EINHEIT_LABEL_ZU_KEY = {label.lower(): key for key, label in EINHEIT_CHOICES}

# Spaltennamen-Präfixe der neun Mindestmenge/VK-Paare, so wie build_interface.py
# (Artikel-Schnittstelle) sie erzeugt: Header "Mindestmenge i\n(Stk)" / "VK i\n(€/Stk)".
# Zeilenumbrüche in Headern werden beim Einlesen normalisiert (siehe _header_name),
# deshalb hier ohne \n gesucht.
STAFFEL_SPALTEN = [(f"Mindestmenge {i} (Stk)", f"VK {i} (€/Stk)") for i in range(1, 10)]


def _s(value):
    if value is None:
        return ""
    return str(value).strip()


def _mehrsprachig(text):
    """Baut aus einer ggf. zweizeiligen Zelle (1. Zeile Deutsch, 2. Zeile Englisch, per
    Zeilenumbruch in derselben Zelle - siehe Docstring oben) einen einzeiligen Text mit
    " // " als Trenner zwischen den Zeilen.

    Hintergrund (Björns Rückmeldung 16.10.2026): der reine Zeilenumbruch kam zwar
    unverändert in Artikel.bezeichnung an, verschwindet aber beim Anzeigen in einzeiligen
    <input>-Feldern (Autofill der Positionszeile) und in der Artikel-Dropdown-Liste
    spurlos - DE und EN kleben dann ohne jede Trennung aneinander (z. B. "... 500er-
    KartonLUAMO Fry Fork ..."). " // " bleibt dagegen in jeder Darstellung sichtbar.
    Zeilen werden einzeln getrimmt, leere Zeilen (z. B. ein überflüssiger Umbruch am
    Ende der Zelle) herausgefiltert; ein Text ohne jeden Zeilenumbruch bleibt unverändert."""
    zeilen = [z.strip() for z in text.splitlines() if z.strip()]
    return " // ".join(zeilen)


def _header_name(value):
    """Normalisiert eine Header-Zelle: Zeilenumbrüche (wrap_text) werden zu Leerzeichen,
    doppelte Leerzeichen zusammengefasst - "Mindestmenge 1\\n(Stk)" wird so zu
    "Mindestmenge 1 (Stk)", unabhängig davon, ob Excel den Umbruch als \\n oder \\r\\n
    gespeichert hat."""
    text = _s(value)
    return " ".join(text.split())


def _dec(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def _int(value):
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


class Command(BaseCommand):
    help = "Importiert den Artikelstamm (Blatt 'Artikel') aus der Artikel-Schnittstelle (Excel)."

    def add_arguments(self, parser):
        parser.add_argument("pfad", type=str, help="Pfad zur .xlsx-Datei")
        parser.add_argument(
            "--dry-run", action="store_true",
            help="Nur anzeigen, was passieren würde - ohne etwas zu speichern.",
        )

    def handle(self, *args, **options):
        pfad = options["pfad"]
        dry_run = options["dry_run"]

        try:
            wb = openpyxl.load_workbook(pfad, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Datei nicht gefunden: {pfad}")

        if "Artikel" not in wb.sheetnames:
            raise CommandError("Blatt 'Artikel' fehlt in der Datei - Import abgebrochen.")

        ws = wb["Artikel"]
        header = [_header_name(c.value) for c in ws[1]]
        spalte = {name: idx for idx, name in enumerate(header) if name}

        # "Bezeichnung *"/"Beschreibung" tragen in der aktuellen Datei den Zusatz
        # "(DE / EN)" im Header - beide Schreibweisen akzeptieren, damit ältere und
        # neuere Dateien funktionieren.
        bezeichnung_spalte = next(
            (n for n in ("Bezeichnung * (DE / EN)", "Bezeichnung *") if n in spalte), None
        )
        beschreibung_spalte = next(
            (n for n in ("Beschreibung (DE / EN)", "Beschreibung") if n in spalte), None
        )

        pflicht = ["Artikelnummer *", "Geschäftsbereich *", "Einheit *", "USt-Satz (%) *"]
        fehlend = [p for p in pflicht if p not in spalte]
        if not bezeichnung_spalte:
            fehlend.append("Bezeichnung *")
        if fehlend:
            raise CommandError(f"Pflichtspalte(n) fehlen im Blatt 'Artikel': {', '.join(fehlend)}")

        einspielfreigabe_spalte = "Einspielfreigabe (Ja/Nein)"
        hat_freigabe_spalte = einspielfreigabe_spalte in spalte

        def wert(row, name):
            idx = spalte.get(name)
            return row[idx] if idx is not None else None

        stats = {"neu": 0, "aktualisiert": 0, "uebersprungen": 0, "nicht_freigegeben": 0, "warnungen": []}

        with transaction.atomic():
            for row in ws.iter_rows(min_row=2, values_only=True):
                artikelnummer = _s(wert(row, "Artikelnummer *"))
                if not artikelnummer:
                    continue

                if hat_freigabe_spalte:
                    freigabe = _s(wert(row, einspielfreigabe_spalte)).lower()
                    if freigabe != "ja":
                        stats["nicht_freigegeben"] += 1
                        continue

                bezeichnung = _mehrsprachig(_s(wert(row, bezeichnung_spalte)))
                if not bezeichnung:
                    stats["warnungen"].append(f"{artikelnummer}: keine Bezeichnung - übersprungen.")
                    stats["uebersprungen"] += 1
                    continue

                warengruppe = _s(wert(row, "Geschäftsbereich *"))
                if not warengruppe:
                    stats["warnungen"].append(f"{artikelnummer}: kein Geschäftsbereich/Warengruppe - übersprungen.")
                    stats["uebersprungen"] += 1
                    continue

                einheit_text = _s(wert(row, "Einheit *")).lower()
                einheit = _EINHEIT_LABEL_ZU_KEY.get(einheit_text)
                if einheit is None:
                    einheit = "stk"
                    stats["warnungen"].append(
                        f"{artikelnummer}: Einheit '{_s(wert(row, 'Einheit *'))}' unbekannt - 'Stück' angenommen."
                    )

                ust_satz = _dec(wert(row, "USt-Satz (%) *"))
                if ust_satz is None:
                    ust_satz = Decimal("19.00")
                    stats["warnungen"].append(f"{artikelnummer}: kein USt-Satz - 19% angenommen.")

                status_text = _s(wert(row, "Status")).lower()
                status = STATUS_MAPPING.get(status_text, "in_vorbereitung")

                defaults = {
                    "bezeichnung": bezeichnung[:300],
                    "beschreibung": _mehrsprachig(_s(wert(row, beschreibung_spalte))) if beschreibung_spalte else "",
                    "geschaeftsbereich": GESCHAEFTSBEREICH_IMPORT,
                    "warengruppe": warengruppe,
                    "einheit": einheit,
                    "verkaufspreis": _dec(wert(row, "Verkaufspreis (EUR, netto) *")),
                    "ust_satz": ust_satz,
                    "einkaufspreis": _dec(wert(row, "Einkaufspreis (EUR, netto)")),
                    "lagerbestand_aktuell": _dec(wert(row, "Lagerbestand aktuell")),
                    "mindestbestand": _dec(wert(row, "Mindestbestand")),
                    "verpackungseinheit_stk": _int(wert(row, "Verpackung Stk")),
                    "lagerort": _s(wert(row, "Lagerort"))[:100],
                    "lieferant": _s(wert(row, "Lieferant"))[:200],
                    "status": status,
                    "notizen": _s(wert(row, "Notizen")),
                }

                obj, created = Artikel.objects.get_or_create(artikelnummer=artikelnummer, defaults=defaults)
                if created:
                    stats["neu"] += 1
                else:
                    for k, v in defaults.items():
                        setattr(obj, k, v)
                    obj.save()
                    stats["aktualisiert"] += 1

                # Preisstaffeln: je Import komplett neu schreiben, direkt aus den
                # Mindestmenge/VK-Spaltenpaaren dieser Zeile (siehe Docstring) - keine
                # warengruppenabhängige Tabelle mehr im Code nötig.
                neue_staffeln = []
                for mindestmenge_spalte, vk_spalte in STAFFEL_SPALTEN:
                    if mindestmenge_spalte not in spalte or vk_spalte not in spalte:
                        continue
                    ab_menge = _int(wert(row, mindestmenge_spalte))
                    preis = _dec(wert(row, vk_spalte))
                    if ab_menge is not None and preis is not None:
                        neue_staffeln.append(Preisstaffel(artikel=obj, ab_menge=ab_menge, preis_netto=preis))
                    elif ab_menge is not None or preis is not None:
                        stats["warnungen"].append(
                            f"{artikelnummer}: {mindestmenge_spalte}/{vk_spalte} nur teilweise befüllt - "
                            f"Stufe übersprungen (beide Zellen müssen gefüllt sein)."
                        )

                if not dry_run:
                    obj.preisstaffeln.all().delete()
                    if neue_staffeln:
                        Preisstaffel.objects.bulk_create(neue_staffeln)

                if defaults["verkaufspreis"] is None and not neue_staffeln:
                    stats["warnungen"].append(f"{artikelnummer}: kein Preis hinterlegt (weder Flatrate noch Staffel).")

            if dry_run:
                transaction.set_rollback(True)
                self.stdout.write(self.style.WARNING("--dry-run: Es wurde NICHTS gespeichert."))

        self.stdout.write(self.style.SUCCESS(
            f"Artikel: {stats['neu']} neu, {stats['aktualisiert']} aktualisiert, "
            f"{stats['uebersprungen']} übersprungen, {stats['nicht_freigegeben']} ohne Einspielfreigabe"
        ))
        if stats["warnungen"]:
            self.stdout.write(self.style.WARNING(f"\n{len(stats['warnungen'])} Warnung(en):"))
            for w in stats["warnungen"]:
                self.stdout.write(self.style.WARNING(f"  - {w}"))
